from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from regulatory_poc.types.models import DocumentChunk, DocumentMetadata


SUPPORTED_EXTENSIONS = {".docx", ".md", ".pdf", ".txt"}


def extract_text(file_name: str, content: bytes) -> str:
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported document type '{extension}'. Use one of: {supported}.")

    if extension in {".txt", ".md"}:
        text = content.decode("utf-8-sig")
    elif extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(
                "PDF ingestion requires pypdf. Run 'python -m pip install --pre -e .'."
            ) from exc
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "DOCX ingestion requires python-docx. Run 'python -m pip install --pre -e .'."
            ) from exc
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        text = "\n".join([*paragraphs, *table_cells])

    normalized = re.sub(r"[ \t]+", " ", text)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if not normalized:
        raise ValueError(
            f"'{file_name}' contained no extractable text. Supply a text-based document or OCR it."
        )
    return normalized


def make_document_id(file_name: str, content: bytes) -> str:
    digest = hashlib.sha256(content).hexdigest()[:16]
    stem = re.sub(r"[^a-z0-9]+", "-", Path(file_name).stem.lower()).strip("-")
    return f"{stem or 'document'}-{digest}"


def chunk_document(
    text: str,
    metadata: DocumentMetadata,
    chunk_words: int = 220,
    overlap_words: int = 40,
) -> list[DocumentChunk]:
    if chunk_words <= 0:
        raise ValueError("chunk_words must be greater than zero.")
    if overlap_words < 0 or overlap_words >= chunk_words:
        raise ValueError("overlap_words must be at least zero and smaller than chunk_words.")

    words = text.split()
    chunks: list[DocumentChunk] = []
    step = chunk_words - overlap_words
    for ordinal, start in enumerate(range(0, len(words), step)):
        chunk_text = " ".join(words[start : start + chunk_words]).strip()
        if not chunk_text:
            continue
        chunk_id = hashlib.sha256(
            f"{metadata.document_id}:{ordinal}:{chunk_text}".encode("utf-8")
        ).hexdigest()
        chunks.append(
            DocumentChunk(
                chunk_id=chunk_id,
                text=chunk_text,
                ordinal=ordinal,
                metadata=metadata,
            )
        )
        if start + chunk_words >= len(words):
            break
    return chunks
