"""Excel export of the comparison table with wrapped text and a frozen header."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_CELL = 32_000
WIDE = 60


def xlsx_table(rows: list[dict], title: str, summary: list[dict] | None = None,
               groups: dict[str, list[str]] | None = None) -> bytes:
    """Comparison sheet first; groups add a merged band above related columns."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Comparison"
    _write_sheet(sheet, rows, wide=WIDE, groups=groups or {})
    if summary is not None:
        _write_sheet(book.create_sheet("Reviewer summary"), summary, wide=80, groups={})
    notes = book.create_sheet("About")
    for line in (
        title,
        "Unreviewed proposed mapping. qualified expert review is required; no compliance certification.",
        "Source quotations are exact extracted text; English translations are machine-generated.",
        "Equivalence is calculated by code from the AI topic checklist; partial topics count as different.",
        "One row per Reference requirement. National clauses are numbered [1], [2]... identically in every column;",
        "only clauses that an reference topic relies on are listed. Proposed analysis is expanded into four columns.",
        "Download the JSON export for lossless strings and full provenance.",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _write_sheet(sheet, rows: list[dict], wide: int, groups: dict[str, list[str]]) -> None:
    headers = list(rows[0]) if rows else []
    bands = [(name, [headers.index(column) + 1 for column in columns if column in headers])
             for name, columns in groups.items()]
    bands = [(name, columns) for name, columns in bands if columns]
    if bands:
        sheet.append([""] * len(headers))
        for name, columns in bands:
            sheet.cell(row=1, column=min(columns), value=name)
            sheet.merge_cells(start_row=1, end_row=1, start_column=min(columns), end_column=max(columns))
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell(row.get(header, "")) for header in headers])
    wrap = Alignment(wrap_text=True, vertical="top")
    for cells in sheet.iter_rows():
        for cell in cells:
            cell.alignment = wrap
            if cell.data_type == "f":
                # Extracted source text is data, never a spreadsheet formula.
                cell.data_type = "s"
    header_row = 2 if bands else 1
    for row_number in range(1, header_row + 1):
        for cell in sheet[row_number]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="BDD7EE" if row_number < header_row else "DDEBF7")
    for index, header in enumerate(headers, start=1):
        longest = max((len(str(row.get(header, ""))) for row in rows), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(wide, longest, 18 + len(header) // 2))
    sheet.freeze_panes = f"B{header_row + 1}"


def _cell(value) -> str:
    text = str(value)
    if len(text) > MAX_CELL:
        text = text[:MAX_CELL] + "\n[truncated in Excel; see JSON export]"
    return text
