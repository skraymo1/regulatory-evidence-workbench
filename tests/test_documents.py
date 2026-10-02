from __future__ import annotations

import io
import unittest

from docx import Document

from regulatory_poc.repo.documents import chunk_document, extract_text, make_document_id
from regulatory_poc.types.models import DocumentMetadata


class DocumentTests(unittest.TestCase):
    def test_extracts_utf8_text_and_builds_stable_id(self) -> None:
        content = "Requirement one.\n\nRequirement two.".encode()
        self.assertEqual(
            extract_text("standard.txt", content),
            "Requirement one.\n\nRequirement two.",
        )
        self.assertEqual(
            make_document_id("standard.txt", content),
            make_document_id("standard.txt", content),
        )

    def test_chunks_with_overlap_and_metadata(self) -> None:
        metadata = DocumentMetadata(
            document_id="doc-1",
            title="Test",
            authority="reference",
        )
        chunks = chunk_document(
            "one two three four five six seven",
            metadata,
            chunk_words=4,
            overlap_words=1,
        )
        self.assertEqual(["one two three four", "four five six seven"], [c.text for c in chunks])
        self.assertTrue(all(chunk.metadata == metadata for chunk in chunks))

    def test_rejects_unsupported_type(self) -> None:
        with self.assertRaisesRegex(ValueError, "Unsupported document type"):
            extract_text("source.html", b"<p>content</p>")

    def test_extracts_docx_table_cells(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Requirement"
        table.cell(0, 1).text = "Defence in depth"
        buffer = io.BytesIO()
        document.save(buffer)

        text = extract_text("regulation.docx", buffer.getvalue())

        self.assertIn("Requirement", text)
        self.assertIn("Defence in depth", text)


if __name__ == "__main__":
    unittest.main()
