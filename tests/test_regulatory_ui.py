from __future__ import annotations

import os
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from streamlit.testing.v1 import AppTest

from regulatory_poc.config.settings import Settings
from regulatory_poc.runtime.regulatory import build_regulatory_service


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "support_docs"


def sample_role(name):
    for prefix, role, language in [
        ("reference", "Reference baseline", "English"), ("UK", "UK", "English"),
        ("Korean Regulations", "Korea", "Korean"),
        ("Regulations on Technical", "Korean English reference", "English"),
        ("Argent", "Argentina", "Spanish"),
    ]:
        if name.startswith(prefix):
            return role, language
    raise ValueError(f"Unknown test fixture: {name}")


def render_saved_table(report):
    from types import SimpleNamespace
    from regulatory_poc.ui.regulatory import render_table
    render_table(SimpleNamespace(agent=None, translation_agent=None), report, "test")


def example_report(country="Argentina"):
    def source(document, identifier, language):
        return {
            "key": document + "-" + identifier, "identifier": identifier, "title": "Original **name**",
            "description": "Exact source line one.\nExact source line two | [not a link].",
            "document_id": document, "source_name": document + ".pdf",
            "source_hash": "a" * 64, "page": 2, "end_page": 3, "language": language,
        }
    baseline = source("reference", "Requirement 1", "English")
    national = source("national", "Article 1", "Korean" if country == "Korea" else "Spanish")
    second = source("second", "Article 2", national["language"])
    metadata = lambda item: {
        "document_id": item["document_id"], "title": item["source_name"],
        "language": item["language"], "revision": "Test revision",
    }
    rows = []
    for index in range(82):
        item = copy.deepcopy(baseline)
        item.update(key=f"reference-{index}", identifier=f"Requirement {index + 1}")
        rows.append({
            "baseline": item, "status": "pending", "matches": [], "candidates": [],
            "note": "", "translation_checks": {},
        })
    rows[0].update(status="proposed", note="Search is not exhaustive.", matches=[
        {"source": item, "english_title": f"Translated name {index}",
         "english_description": "Machine translation, not a quotation.",
         "analysis": "Proposed relationship requires expert review."}
        for index, item in enumerate([national, second], start=1)
    ])
    rows[1].update(status="unresolved", note="No candidate found; not proof of absence.")
    rows[2].update(status="error", note="Source retrieval failed. Retry the row.")
    rows[3].update(status="candidates_only", candidates=[national], note="No technical mapping in offline mode.")
    return {
        "report_id": "a" * 64, "inputs": {
            "country": country, "baseline": metadata(baseline),
            "sources": [metadata(national), metadata(second)],
            "reference_id": "", "selected_ids": [],
        },
        "rows": rows, "warnings": ["Provisional extraction; review the originals."],
        "review_status": "unreviewed", "indexing_status": "not_started",
        "inventories": {"national": [national], "second": [second]},
    }


class RegulatoryUiTests(unittest.TestCase):
    def test_annex_detail_precedes_complete_table_and_keeps_exact_source_text(self):
        report = example_report()
        app = AppTest.from_function(render_saved_table, args=(report,), default_timeout=30).run()
        self.assertFalse(app.exception)
        self.assertEqual(["Comparison table", "Reviewer summary (3 columns)", "Requirement detail"],
                         [tab.label for tab in app.tabs])
        detail = app.tabs[2]
        self.assertEqual(82, len(detail.selectbox[0].options))
        self.assertFalse(detail.dataframe)
        self.assertFalse(detail.json)
        code = [block.value for block in detail.code]
        self.assertEqual(report["rows"][0]["baseline"]["title"], code[0])
        self.assertEqual(report["rows"][0]["baseline"]["description"], code[1])
        self.assertEqual("Translated name 1", code[4])
        self.assertEqual(report["rows"][0]["matches"][1]["source"]["title"], code[6])
        self.assertEqual("Translated name 2", code[8])
        table = app.tabs[0].dataframe[0].value
        self.assertEqual(82, len(table))
        first = table.iloc[0]
        self.assertTrue(first["National requirement(s)"].startswith("[1] "))
        self.assertIn("\n\n[2] second, Article 2", first["National requirement(s)"])
        self.assertIn("[2] Proposed relationship requires expert review.", first["Clause-by-clause analysis (English)"])
        self.assertEqual("Search is not exhaustive.", first["Evidence limitations"])
        self.assertIn("second.pdf", table.iloc[0]["National source reference(s)"])
        self.assertIn("PDF pages 2-3", table.iloc[0]["Reference source reference"])
        self.assertIn("Translated name 2", table.iloc[0]["National English translation (machine-generated)"])
        self.assertEqual("Not assessed - Proposed correspondence - expert review required",
                         table.iloc[0]["Are the requirements equivalent?"])
        self.assertNotIn("Proposed analysis (English)", table.columns)
        columns = list(table.columns)
        self.assertEqual(
            ["National source reference(s)", "Are the requirements equivalent?"],
            columns[columns.index("National source reference(s)"):][:2])
        self.assertEqual(columns.index("Clause-by-clause analysis (English)") + 1,
                         columns.index("Evidence limitations"))
        metrics = {metric.label: metric.value for metric in app.metric}
        self.assertEqual({"Reference requirements": "82", "Proposed mappings": "1",
                          "Unresolved / candidates only": "2", "Pending / retry": "79"}, metrics)
        for index in (1, 2, 3, 81):
            app.tabs[2].selectbox[0].select(f"reference-{index}").run()
            self.assertFalse(app.exception)
            self.assertTrue(any(report["rows"][index]["note"] in info.value
                                for info in app.tabs[2].info))
            self.assertFalse(any("National content" in text.value for text in app.tabs[2].markdown))

    def test_three_column_equivalence_format_detail_and_excel_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from regulatory_poc.ui.regulatory import SUMMARY_COLUMNS, summary_rows, table_rows
        from regulatory_poc.ui.regulatory_columns import (
            ANALYSIS_COLUMNS, EQUIVALENCE, NATIONAL_ONLY, PROPOSED_ANALYSIS, REFERENCE_ONLY, TERMINOLOGY,
        )
        from regulatory_poc.ui.regulatory_export import xlsx_table
        report = example_report()
        national, second = (match["source"] for match in report["rows"][0]["matches"])
        report["rows"][0].update(
            paragraphs=[{"identifier": "6.20", "text": "6.20. Exact supporting paragraph.", "page": 5,
                         "end_page": 5, "removed": []}],
            candidates=[national, second], equivalence="small_differences",
            topics=[
                {"id": "T1", "topic": "Leak tightness", "source": "statement", "status": "explicit",
                 "keys": [national["key"]], "explanation": ""},
                {"id": "T2", "topic": "Aircraft crash", "source": "6.20", "status": "not_explicit",
                 "keys": [], "explanation": ""},
                {"id": "T3", "topic": "Isolation", "source": "6.21", "status": "explicit",
                 "keys": [second["key"]], "explanation": ""},
            ],
            national_only_topics=[{"topic": "Periodic reporting", "keys": [second["key"]]}],
            terminology=[{"term": "confinamiento", "language": "Spanish", "literal_english": "confinement",
                          "reference_term": "containment", "concern": "system vs safety function"}],
            cited_instruments=["AR 10.10.1"],
            screening={"complete": True, "clauses": 157, "batches": 1, "passes": [{}, {}]},
        )
        row = table_rows(report)[0]
        self.assertTrue(row["Are the requirements equivalent?"].startswith(
            "Small number of differences (1 of 3 reference topics differ"))
        self.assertIn("[6.20] Aircraft crash (not explicit)", row[REFERENCE_ONLY])
        self.assertNotIn("Leak tightness", row[REFERENCE_ONLY])
        self.assertIn("Periodic reporting (in [2])", row[NATIONAL_ONLY])
        self.assertIn("confinamiento", row[TERMINOLOGY])
        self.assertEqual("6.20. Exact supporting paragraph.", row["Reference exact supporting paragraphs"])
        sheet = load_workbook(BytesIO(xlsx_table(table_rows(report), "Test"))).active
        headers = [cell.value for cell in sheet[1]]
        self.assertIn(REFERENCE_ONLY, headers)
        self.assertEqual(len(table_rows(report)), sheet.max_row - 1)
        grouped = load_workbook(BytesIO(xlsx_table(
            table_rows(report), "Test", None, {PROPOSED_ANALYSIS: list(ANALYSIS_COLUMNS)}))).active
        headers = [cell.value for cell in grouped[2]]
        start = headers.index(EQUIVALENCE) + 1
        self.assertEqual(PROPOSED_ANALYSIS, grouped.cell(row=1, column=start).value)
        self.assertIn(f"{grouped.cell(row=1, column=start).coordinate}:"
                      f"{grouped.cell(row=1, column=start + 3).coordinate}",
                      [str(item) for item in grouped.merged_cells.ranges])
        self.assertEqual(len(table_rows(report)), grouped.max_row - 2)
        summary = summary_rows(report)
        self.assertEqual(list(SUMMARY_COLUMNS), list(summary[0]))
        self.assertEqual(3, len([column for column in summary[0]
                                 if column in (EQUIVALENCE, REFERENCE_ONLY, NATIONAL_ONLY)]))
        self.assertEqual(row[REFERENCE_ONLY], summary[0][REFERENCE_ONLY])
        self.assertNotIn("\n\n", summary[0]["National requirement(s)"])
        book = load_workbook(BytesIO(xlsx_table(table_rows(report), "Test", summary)))
        self.assertEqual(["Comparison", "Reviewer summary", "About"], book.sheetnames)
        self.assertEqual(list(SUMMARY_COLUMNS), [cell.value for cell in book["Reviewer summary"][1]])
        self.assertEqual(len(summary), book["Reviewer summary"].max_row - 1)
        app = AppTest.from_function(render_saved_table, args=(report,), default_timeout=30).run()
        self.assertFalse(app.exception)
        self.assertEqual(list(SUMMARY_COLUMNS), list(app.tabs[1].dataframe[0].value.columns))
        detail = app.tabs[2]
        self.assertTrue(any("Proposed equivalence assessment" in item.value for item in detail.markdown))
        self.assertIn("6.20. Exact supporting paragraph.", [item.value for item in detail.code])
        self.assertTrue(any("157 national clauses" in item.value and "2 pass(es)" in item.value
                            for item in detail.caption))

    def test_korean_reference_remains_separate_and_verification_available_before_mapping(self):
        report = example_report("Korea")
        korean = report["inventories"]["national"][0]
        reference = copy.deepcopy(korean)
        reference.update(key="english-1", document_id="english", source_name="english.pdf",
                         title="Supplied English name", description="Supplied English description.",
                         language="English")
        report["inputs"].update(reference_id="english", selected_ids=["Article 1"])
        report["inputs"]["sources"] = [
            report["inputs"]["sources"][0],
            {"document_id": "english", "title": "English reference", "language": "English", "revision": "Older"},
        ]
        report["inventories"]["english"] = [reference]
        report["rows"][0].update(status="pending", matches=[])
        app = AppTest.from_function(render_saved_table, args=(report,), default_timeout=30).run()
        self.assertFalse(app.exception)
        self.assertTrue(any(box.label == "Original requirement to verify" for box in app.selectbox))
        self.assertTrue(any(button.label == "Verify this requirement's translation" for button in app.button))
        self.assertEqual("Translation reference", app.dataframe[0].value.iloc[2]["Role"])
        korean["english_reference"] = reference
        report["rows"][0].update(status="proposed", matches=[{
            "source": korean, "analysis": "Unreviewed technical analysis.",
            "english_title": "Generated English name", "english_description": "Generated English description.",
        }])
        app = AppTest.from_function(render_saved_table, args=(report,), default_timeout=30).run()
        self.assertFalse(app.exception)
        code = [item.value for item in app.tabs[2].code]
        self.assertLess(code.index("Generated English name"), code.index("Supplied English name"))
        self.assertTrue(any("not a certified translation" in item.value for item in app.tabs[2].caption))

    @unittest.skipUnless(SAMPLES.is_dir(), "Local support_docs PDFs not present")
    def test_real_sources_three_tables_and_resumable_ui(self):
        home = str(Path.home())
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {
            "POC_AGENT_MODE": "offline", "POC_SEARCH_MODE": "local",
            "POC_INDEX_PATH": str(Path(directory) / "index.json"),
            "USERPROFILE": home, "HOME": home,
        }, clear=True):
            service = build_regulatory_service(Settings.from_env())
            for path in sorted(SAMPLES.glob("*.pdf")):
                role, language = sample_role(path.name)
                service.library.import_document(
                    path.read_bytes(), path.name, role, language,
                    "Reference standard" if role == "Reference baseline" else "",
                    "Rev. 1" if role == "Reference baseline" else "",
                )
            records = service.library.list_documents()
            self.assertEqual(15, len(records))
            self.assertTrue(all(item["indexing_status"] == "indexed" for item in records))
            baseline = next(item["metadata"]["document_id"] for item in records if item["role"] == "Reference baseline")
            for country in ("UK", "Korea", "Argentina"):
                sources = [item["metadata"]["document_id"] for item in records if item["role"] == country]
                selection = [service.library.requirements(sources[0])[0].identifier] if country == "Korea" else []
                report = service.create(baseline, country, sources, selection)
                self.assertEqual(82, len(report["rows"]))
                if country == "Argentina":
                    self.assertEqual(11, len(report["inputs"]["sources"]))
            app = AppTest.from_file(ROOT / "src" / "regulatory_poc" / "ui" / "app.py", default_timeout=90).run()
            self.assertFalse(app.exception)
            self.assertEqual("Comparisons", app.tabs[0].label)
            labels = {tab.label for tab in app.tabs}
            self.assertTrue({"Comparison sources", "Saved comparison tables", "Ask evidence"} <= labels)
            self.assertTrue({"Translation fidelity", "Upload", "Repositories", "Saved reports", "Chat"}.isdisjoint(labels))
            self.assertFalse(app.sidebar.metric)
            self.assertFalse(app.sidebar.caption)
            self.assertFalse(app.json)
            self.assertFalse(any(button.label == "Verify and index reference samples" for button in app.button))
            self.assertEqual("Review saved table", app.radio(key="comparison-workspace").value)
            self.assertTrue(any(box.label == "Saved comparison" for box in app.selectbox))
            app.radio(key="comparison-workspace").set_value("Set up a comparison").run()
            app.selectbox(key="comparison-jurisdiction").select("UK").run()
            app.checkbox(key="comparison-scope-acknowledged").check().run()
            app.selectbox(key="create-process-count").select(1).run()
            app.button(key="create-comparison").click().run()
            self.assertFalse(app.exception)
            matching = [table for table in app.dataframe if "Row type" in table.value.columns]
            self.assertTrue(matching)
            expected_count = 82 + len(service.library.requirements(
                next(item["metadata"]["document_id"] for item in records if item["role"] == "UK")
            ))
            self.assertEqual(expected_count, len(matching[0].value))
            self.assertEqual(82, sum(matching[0].value["Row type"] == "Reference baseline"))
            detail = next(box for box in app.selectbox if box.label == "Open requirement detail (Annex A view)")
            self.assertEqual(expected_count, len(detail.options))
            app.radio(key="comparison-workspace").set_value("Review saved table").run()
            uk_report = next(item for item in service.list_reports() if item["inputs"]["country"] == "UK")
            next(box for box in app.selectbox if box.label == "Saved comparison").select(uk_report["report_id"]).run()
            next(box for box in app.selectbox if box.label == "Rows to process now").select(1).run()
            next(button for button in app.button
                 if button.label == "Process pending requirements / retry errors").click().run()
            self.assertFalse(app.exception)
            matching = [table for table in app.dataframe if "Row type" in table.value.columns]
            self.assertEqual("candidates_only", matching[0].value.iloc[0]["Status"])
            self.assertTrue(any(button.label == "Download comparison table (Markdown)"
                                for button in app.get("download_button")))
